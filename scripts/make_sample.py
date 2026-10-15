"""Generate the sample ML course (our own content) under data/sample_courses/ml/.

Produces:
  lecture_slides.pptx            6 slides; slide 6 carries a chart (learning-rate loss curves)
  lecture_transcript.vtt         timestamped transcript of lecture one
  model_selection_chapter.pdf    3-page textbook chapter with two embedded figures
  lecture2_model_selection.mp4   narrated lecture (TTS over still frames) used to demo Whisper ingestion
ml_notes.md and topic_map.json are authored by hand and not touched here.

Needs: python-pptx, pymupdf, matplotlib, edge-tts (online; falls back to Windows speech), imageio-ffmpeg.
Usage: python scripts/make_sample.py [--no-video]
"""
from __future__ import annotations

import asyncio
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).parent.parent / "data" / "sample_courses" / "ml"
TMP = Path(tempfile.mkdtemp(prefix="sc_sample_"))

# ------------------------------------------------------------------ figures

def fig_learning_rates() -> Path:
    it = np.arange(0, 41)
    fig, ax = plt.subplots(figsize=(7, 4.2), dpi=110)
    ax.plot(it, 8 * np.exp(-0.05 * it) + 0.4, label="α = 0.01 (too small)", color="#1f77b4")
    ax.plot(it, 8 * np.exp(-0.28 * it) + 0.4, label="α = 0.1 (good)", color="#2ca02c")
    ax.plot(it, 1.0 + 0.9 * 1.18 ** it, label="α = 1.2 (too large: diverges)", color="#d62728")
    ax.set_ylim(0, 12)
    ax.set_xlabel("Iteration")
    ax.set_ylabel("Training loss")
    ax.set_title("Loss during gradient descent for three learning rates")
    ax.legend()
    ax.grid(alpha=0.3)
    p = TMP / "fig_learning_rates.png"
    fig.tight_layout()
    fig.savefig(p)
    plt.close(fig)
    return p


def fig_bias_variance() -> Path:
    x = np.linspace(1, 10, 100)
    train = 6.5 * np.exp(-0.38 * x) + 0.3
    val = 6.5 * np.exp(-0.38 * x) + 0.3 + 0.012 * (x - 1) ** 2.4 * 0.6
    best = x[np.argmin(val)]
    fig, ax = plt.subplots(figsize=(7, 4.2), dpi=110)
    ax.plot(x, train, label="Training error", color="#1f77b4")
    ax.plot(x, val, label="Validation error", color="#d62728")
    ax.axvline(best, ls="--", color="gray")
    ax.text(best + 0.1, val.max() * 0.8, f"best complexity ≈ {best:.0f}", color="gray")
    ax.text(1.2, val.max() * 0.55, "underfitting\n(high bias)", color="#444")
    ax.text(7.4, val.max() * 0.35, "overfitting\n(high variance)", color="#444")
    ax.set_xlabel("Model complexity (polynomial degree)")
    ax.set_ylabel("Error")
    ax.set_title("Training vs validation error as complexity grows")
    ax.legend(loc="upper center")
    p = TMP / "fig_bias_variance.png"
    fig.tight_layout()
    fig.savefig(p)
    plt.close(fig)
    return p


def fig_kfold() -> Path:
    fig, ax = plt.subplots(figsize=(7, 3.6), dpi=110)
    k = 5
    for row in range(k):
        for col in range(k):
            val = row == col
            ax.add_patch(plt.Rectangle((col, k - row - 1), 0.92, 0.8, color="#ff7f0e" if val else "#9ecae1"))
            ax.text(col + 0.46, k - row - 0.6, "validation" if val else "train", ha="center", va="center", fontsize=8)
        ax.text(-0.15, k - row - 0.6, f"Iteration {row + 1}", ha="right", va="center", fontsize=9)
    for col in range(k):
        ax.text(col + 0.46, k + 0.1, f"Fold {col + 1}", ha="center", fontsize=9)
    ax.set_xlim(-1.6, k)
    ax.set_ylim(-0.2, k + 0.6)
    ax.axis("off")
    ax.set_title("5-fold cross-validation: every fold is the validation set once")
    p = TMP / "fig_kfold.png"
    fig.tight_layout()
    fig.savefig(p)
    plt.close(fig)
    return p


# ------------------------------------------------------------------- slides

SLIDES = [
    ("Introduction to Machine Learning", "Lecture 1 — what learning from data means", [
        "Machine learning: algorithms that improve with experience",
        "Task T, experience E, performance measure P",
        "Three paradigms: supervised, unsupervised, reinforcement",
    ], None),
    ("Supervised vs Unsupervised", "The two workhorses", [
        "Supervised: labelled (x, y) pairs — classification & regression",
        "Unsupervised: unlabelled data — clustering, dimensionality reduction",
        "Reinforcement: agent, environment, reward signal",
    ], None),
    ("Linear Regression", "Fitting a line to data", [
        "Prediction: y_hat = w·x + b",
        "Loss: mean squared error (MSE)",
        "Closed form: normal equations w = (X^T X)^{-1} X^T y",
        "O(d^3) inversion — gradient descent often preferred",
    ], None),
    ("Gradient Descent", "The optimization engine", [
        "theta := theta - alpha * grad L(theta)",
        "Learning rate alpha: too small = slow, too large = diverge",
        "SGD uses mini-batches for cheaper steps",
    ], None),
    ("Overfitting & Regularization", "When memorizing beats learning", [
        "Overfitting: great on train, poor on new data",
        "L2 (ridge): lambda * ||w||^2 — smooth shrinkage",
        "L1 (lasso): lambda * ||w||_1 — sparse, feature selection",
        "Tune lambda on validation, report on test",
    ], None),
    ("Choosing the Learning Rate", "Reading the loss curve", [
        "Plot the training loss every iteration",
        "Compare three step sizes side by side",
    ], "lr"),
]


def make_slides(lr_fig: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    blank = prs.slide_layouts[6]
    for title, subtitle, bullets, fig in SLIDES:
        slide = prs.slides.add_slide(blank)
        tf = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(1.4)).text_frame
        tf.text = title
        tf.paragraphs[0].font.size = Pt(34)
        tf.paragraphs[0].font.bold = True
        p = tf.add_paragraph()
        p.text = subtitle
        p.font.size = Pt(17)
        width = 5.8 if fig else 11.5
        bf = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(width), Inches(4.5)).text_frame
        bf.word_wrap = True
        for i, b in enumerate(bullets):
            para = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
            para.text = "•  " + b
            para.font.size = Pt(20)
            para.space_after = Pt(14)
        if fig:
            slide.shapes.add_picture(str(lr_fig), Inches(6.9), Inches(1.9), width=Inches(5.9))
        slide.notes_slide.notes_text_frame.text = (
            "Explain: " + " / ".join(bullets) if not fig else "Walk through the figure with the class."
        )
    prs.save(OUT / "lecture_slides.pptx")
    print("slides:", len(SLIDES))


# ---------------------------------------------------------------------- PDF

def make_pdf(bv: Path, kf: Path) -> None:
    import fitz

    doc = fitz.open()

    def page(title: str | None = None):
        pg = doc.new_page(width=595, height=842)
        if title:
            pg.insert_text((56, 70), title, fontsize=20, fontname="hebo")
        return pg

    def para(pg, y, text, h=120, size=11):
        pg.insert_textbox(fitz.Rect(56, y, 539, y + h), text, fontsize=size, fontname="helv", align=0)

    p1 = page("Chapter 5: Model Selection")
    para(p1, 100, "5.1 Bias and variance", 24, 14)
    para(p1, 128,
         "A model with high bias is too simple to capture the pattern in the data. It underfits: both its training error "
         "and its validation error are high. A model with high variance is too flexible. It overfits: its training error is "
         "very low but its validation error is high, because it has memorized noise in the training set.", 110)
    para(p1, 245,
         "As model complexity increases, the training error keeps decreasing. The validation error first decreases and then "
         "increases again. The best model is the one near the minimum of the validation error curve. Figure 5.1 on the next "
         "page shows this pattern for polynomial regression.", 100)

    p2 = page()
    p2.insert_image(fitz.Rect(56, 60, 539, 345), filename=str(bv))
    para(p2, 352, "Figure 5.1: Training and validation error as model complexity grows. Training error falls steadily; "
                  "validation error is U-shaped.", 40, 10)
    para(p2, 410,
         "The gap between training and validation error is a direct measure of overfitting. When the gap is wide, "
         "reduce model complexity, add regularization, or collect more data. Never choose the complexity using the test set.", 90)

    p3 = page()
    para(p3, 70, "5.2 Cross-validation", 24, 14)
    para(p3, 98,
         "When data is limited, a single validation split gives a noisy estimate. k-fold cross-validation splits the data "
         "into k equal folds. Each fold is used once as the validation set while the other k-1 folds train the model. "
         "The k validation scores are averaged to estimate performance. Typical choices are k = 5 or k = 10; "
         "leave-one-out is the extreme case k = n.", 120)
    p3.insert_image(fitz.Rect(56, 235, 539, 235 + 483 * 3.6 / 7), filename=str(kf))
    para(p3, 235 + 483 * 3.6 / 7 + 8,
         "Figure 5.2: 5-fold cross-validation. In each iteration a different fold is held out for validation.", 40, 10)
    doc.save(OUT / "model_selection_chapter.pdf")
    print("pdf pages:", doc.page_count)


# -------------------------------------------------------------------- video

LECTURE2 = [
    ("Lecture two: model selection. A model with high bias is too simple to capture the pattern in the data. It "
     "underfits, so both its training error and its validation error stay high. A model with high variance is too "
     "flexible. It overfits, so its training error is very low but its validation error is high, because it has "
     "memorized noise. As model complexity grows, the training error keeps falling, while the validation error first "
     "falls and then rises again. We pick the complexity near the minimum of the validation curve, and we never use "
     "the test set to make that choice.", "bv"),
    ("When data is limited, a single validation split gives a noisy estimate, so we use k-fold cross-validation. Split "
     "the data into k equal folds. Each fold is used once as the validation set, while the other folds are used for "
     "training. Then average the k validation scores to estimate how well the model will perform. Typical choices are "
     "k equals five or k equals ten, and leave-one-out is the extreme case where k equals the number of examples.", "kf"),
]


async def _edge_tts(text: str, out: Path) -> None:
    import edge_tts

    await edge_tts.Communicate(text, "en-US-AriaNeural", rate="-5%").save(str(out))


def _tts(text: str, out: Path) -> None:
    try:
        asyncio.run(_edge_tts(text, out))
    except Exception:  # offline: Windows built-in speech -> wav
        out = out.with_suffix(".wav")
        ps = (f"Add-Type -AssemblyName System.Speech; $s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
              f"$s.SetOutputToWaveFile('{out}'); $s.Speak('{text.replace(chr(39), '')}'); $s.Dispose()")
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


def _frame(title: str, fig: Path, out: Path) -> None:
    from PIL import Image

    f = plt.figure(figsize=(12.8, 7.2), dpi=100)
    f.patch.set_facecolor("white")
    f.text(0.05, 0.9, title, fontsize=30, weight="bold")
    f.text(0.05, 0.83, "Lecture 2 — Model selection", fontsize=16, color="#555")
    img = Image.open(fig)
    ax = f.add_axes([0.14, 0.04, 0.72, 0.74])
    ax.imshow(img)
    ax.axis("off")
    f.savefig(out)
    plt.close(f)


def make_video(bv: Path, kf: Path) -> None:
    import imageio_ffmpeg

    ff = imageio_ffmpeg.get_ffmpeg_exe()
    clips = []
    for i, (text, which) in enumerate(LECTURE2):
        audio = TMP / f"narr{i}.mp3"
        _tts(text, audio)
        if not audio.exists():
            audio = audio.with_suffix(".wav")
        frame = TMP / f"frame{i}.png"
        _frame("Bias, variance and model complexity" if which == "bv" else "k-fold cross-validation", bv if which == "bv" else kf, frame)
        clip = TMP / f"clip{i}.mp4"
        subprocess.run([ff, "-y", "-loglevel", "error", "-loop", "1", "-i", str(frame), "-i", str(audio), "-c:v", "libx264",
                        "-tune", "stillimage", "-pix_fmt", "yuv420p", "-r", "10", "-c:a", "aac", "-b:a", "128k",
                        "-shortest", str(clip)], check=True)
        clips.append(clip)
    lst = TMP / "list.txt"
    lst.write_text("".join(f"file '{c.as_posix()}'\n" for c in clips), encoding="utf-8")
    subprocess.run([ff, "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", str(lst), "-c", "copy",
                    str(OUT / "lecture2_model_selection.mp4")], check=True)
    print("video written")


# ---------------------------------------------------------------------- VTT

VTT = """WEBVTT

00:00:00.000 --> 00:00:22.000
Welcome to lecture one. Machine learning is the study of algorithms that improve their performance on a task through experience. Formally, a program learns from experience E with respect to task T and performance measure P if its performance at T improves with E.

00:00:22.000 --> 00:00:48.000
There are three main paradigms. Supervised learning uses labelled examples, input-output pairs, and learns a mapping that generalizes to unseen data. Think spam classification or house price prediction.

00:00:48.000 --> 00:01:15.000
Unsupervised learning works with unlabelled data and tries to discover structure: clusters, low-dimensional representations, or anomalies. Customer segmentation is the classic example.

00:01:15.000 --> 00:01:40.000
Reinforcement learning trains an agent that interacts with an environment. The agent observes states, takes actions, and receives rewards, and the goal is a policy that maximizes expected cumulative reward.

00:01:40.000 --> 00:02:10.000
Now linear regression. We model a continuous target y as y_hat equals w dot x plus b. Training means picking w and b to minimize a loss, most commonly the mean squared error. For MSE there is a closed-form solution, the normal equations: w equals X transpose X inverse, times X transpose y.

00:02:10.000 --> 00:02:40.000
But inverting X transpose X costs order d cubed, which is painful when you have many features, and it can be numerically unstable. That is why in practice we usually prefer gradient descent.

00:02:40.000 --> 00:03:15.000
Gradient descent is iterative. You start from an initial guess and repeatedly step in the direction of the negative gradient: theta becomes theta minus alpha times the gradient of the loss. Alpha is the learning rate and it is the most important hyperparameter. Too small and training crawls. Too large and you overshoot and diverge.

00:03:15.000 --> 00:03:45.000
Stochastic gradient descent approximates the gradient with a small random mini-batch. Each step is much cheaper, and convergence is often faster in practice, but the loss curve gets noisier.

00:03:45.000 --> 00:04:20.000
Finally, overfitting and regularization. A model overfits when it fits the training data very well but generalizes poorly. The remedy is to add a penalty on complexity. L2, or ridge, adds lambda times the squared norm of w, shrinking weights smoothly. L1, or lasso, adds lambda times the absolute norm, which pushes some weights exactly to zero and performs feature selection. Always tune lambda on a validation set, and report final numbers on the test set.
"""


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    lr, bv, kf = fig_learning_rates(), fig_bias_variance(), fig_kfold()
    make_slides(lr)
    (OUT / "lecture_transcript.vtt").write_text(VTT, encoding="utf-8")
    make_pdf(bv, kf)
    if "--no-video" not in sys.argv:
        make_video(bv, kf)
    shutil.rmtree(TMP, ignore_errors=True)


if __name__ == "__main__":
    main()
