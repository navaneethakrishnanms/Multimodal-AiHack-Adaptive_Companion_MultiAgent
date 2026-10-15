"""Multimodal ingestion — turn course material into source-linked content units.

Supported inputs:
  - PDF  (textbook chapters, papers, notes)   -> pymupdf, per-page text + embedded image refs
  - PPTX (lecture slide decks)                -> python-pptx, per-slide text + speaker notes
  - TXT/MD                                    -> paragraph chunks
  - VTT/SRT/TXT transcripts (lecture videos)  -> timestamped chunks

Every unit keeps a precise origin: file, page / slide / timestamp.
"""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, asdict
from pathlib import Path


@dataclass
class Unit:
    id: str
    text: str
    source_type: str  # pdf | pptx | text | transcript | video | audio
    source_file: str
    location: str  # "p.12" | "slide 4" | "00:12:30" | "para 3"
    page: int | None = None
    image_count: int = 0
    topics: list[str] | None = None
    kind: str = "text"  # text | figure
    figure_path: str | None = None  # data/figures/<file> for figure units
    start_sec: float | None = None  # video/audio/transcript start, for seeking
    section: str | None = None  # nearest heading (markdown notes)

    def to_dict(self) -> dict:
        return asdict(self)


def _uid(file: str, loc: str, idx: int) -> str:
    h = hashlib.sha1(f"{file}:{loc}:{idx}".encode()).hexdigest()[:8]
    return f"u{h}"


def _chunk_paragraphs(text: str, max_chars: int = 900) -> list[str]:
    """Split text into paragraph-ish chunks under max_chars."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        if len(buf) + len(p) + 2 <= max_chars:
            buf = f"{buf}\n\n{p}" if buf else p
        else:
            if buf:
                chunks.append(buf)
            # hard-split very long paragraphs
            while len(p) > max_chars:
                cut = p.rfind(". ", 0, max_chars)
                cut = cut + 1 if cut > max_chars // 2 else max_chars
                chunks.append(p[:cut].strip())
                p = p[cut:].strip()
            buf = p
    if buf:
        chunks.append(buf)
    return chunks


def _figure_unit(path: Path, loc: str, fig_no: int, blob: bytes, ext: str, nearby: str,
                 figure_dir: Path, source_type: str, page: int | None = None) -> Unit | None:
    """Save one extracted image and turn it into a citable figure unit (described if a vision model is up)."""
    from . import vision

    img = vision.save_image(blob, ext, figure_dir)
    if img is None:
        return None
    desc, _mode = vision.describe(img, nearby)
    return Unit(
        id=_uid(path.name, f"{loc}-fig", fig_no),
        text=vision.figure_text(loc, desc, nearby),
        source_type=source_type,
        source_file=path.name,
        location=f"{loc} · figure {fig_no}",
        page=page,
        kind="figure",
        figure_path=img.name,
        image_count=1,
    )


def ingest_pdf(path: Path, figure_dir: Path | None = None) -> list[Unit]:
    import fitz  # pymupdf

    units: list[Unit] = []
    doc = fitz.open(path)
    for pno, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        n_images = len(page.get_images(full=True))
        if figure_dir is not None:
            for fi, info in enumerate(page.get_images(full=True), start=1):
                try:
                    xref = info[0]
                    img = doc.extract_image(xref)
                    rects = page.get_image_rects(xref)
                    nearby = ""
                    if rects:
                        r = rects[0]
                        nearby = page.get_text("text", clip=fitz.Rect(r.x0 - 10, r.y1, r.x1 + 10, r.y1 + 60)).strip()
                    fu = _figure_unit(path, f"p.{pno}", fi, img["image"], img["ext"], nearby, figure_dir, "pdf", pno)
                except Exception:
                    fu = None
                if fu:
                    units.append(fu)
        if not text and n_images == 0:
            continue
        for ci, chunk in enumerate(_chunk_paragraphs(text)):
            if len(chunk) < 40:
                continue
            units.append(
                Unit(
                    id=_uid(path.name, f"p{pno}", ci),
                    text=chunk,
                    source_type="pdf",
                    source_file=path.name,
                    location=f"p.{pno}",
                    page=pno,
                    image_count=n_images if ci == 0 else 0,
                )
            )
    doc.close()
    return units


def ingest_pptx(path: Path, figure_dir: Path | None = None) -> list[Unit]:
    from pptx import Presentation

    units: list[Unit] = []
    prs = Presentation(str(path))
    for sno, slide in enumerate(prs.slides, start=1):
        parts = []
        n_images = 0
        pictures = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
            if shape.shape_type == 13:  # PICTURE
                n_images += 1
                pictures.append(shape)
        if figure_dir is not None:
            title = parts[0].splitlines()[0] if parts else ""
            for fi, shape in enumerate(pictures, start=1):
                try:
                    fu = _figure_unit(path, f"slide {sno}", fi, shape.image.blob, shape.image.ext, title, figure_dir, "pptx")
                except Exception:
                    fu = None
                if fu:
                    units.append(fu)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(parts)
        if notes:
            body += f"\n\n[Speaker notes] {notes}"
        if not body.strip():
            continue
        units.append(
            Unit(
                id=_uid(path.name, f"slide{sno}", 0),
                text=body.strip(),
                source_type="pptx",
                source_file=path.name,
                location=f"slide {sno}",
                image_count=n_images,
            )
        )
    return units


def ingest_text(path: Path) -> list[Unit]:
    """Notes / markdown: split at headings first so a chunk never straddles two concepts."""
    text = path.read_text(encoding="utf-8", errors="replace")
    sections: list[tuple[str | None, str]] = []
    heading, buf = None, []
    for line in text.splitlines():
        if re.match(r"^#{1,4}\s", line):
            if "".join(buf).strip():
                sections.append((heading, "\n".join(buf)))
            heading, buf = line.lstrip("# ").strip(), [line]
        else:
            buf.append(line)
    if "".join(buf).strip():
        sections.append((heading, "\n".join(buf)))
    units, ci = [], 0
    for heading, body in sections:
        for chunk in _chunk_paragraphs(body, max_chars=700):
            if len(chunk) < 40 or (heading and chunk.strip().lstrip("# ").strip() == heading):
                continue
            ci += 1
            units.append(
                Unit(
                    id=_uid(path.name, "text", ci),
                    text=chunk,
                    source_type="text",
                    source_file=path.name,
                    location=f"para {ci}",
                    section=heading,
                )
            )
    return units


_TS_RE = re.compile(
    r"(?:(\d{1,2}):)?(\d{1,2}):(\d{2})[.,]\d{3}\s*-->"
)


def _ts_to_sec(ts: str) -> int:
    parts = ts.replace(",", ".").split(":")
    parts = [float(p) for p in parts]
    if len(parts) == 3:
        return int(parts[0] * 3600 + parts[1] * 60 + parts[2])
    if len(parts) == 2:
        return int(parts[0] * 60 + parts[1])
    return int(parts[0])


def _hms(sec: float) -> str:
    sec = int(sec)
    mm, ss = divmod(sec, 60)
    hh, mm = divmod(mm, 60)
    return f"{hh:02d}:{mm:02d}:{ss:02d}"


def _windows(entries: list[tuple[float, str]], name: str, source_type: str, window: float = 60.0) -> list[Unit]:
    """Group (start_seconds, text) cues into ~`window`-second units that remember where they start."""
    units: list[Unit] = []
    buf: list[str] = []
    start = entries[0][0]

    def flush() -> None:
        if buf:
            units.append(
                Unit(
                    id=_uid(name, f"t{int(start)}", 0),
                    text=" ".join(buf).strip(),
                    source_type=source_type,
                    source_file=name,
                    location=_hms(start),
                    start_sec=float(int(start)),
                )
            )

    for ts, text in entries:
        if buf and ts - start >= window:  # close the window BEFORE this cue so the label matches the first cue
            flush()
            buf, start = [], ts
        buf.append(text)
    flush()
    return units


def ingest_transcript(path: Path) -> list[Unit]:
    """VTT or SRT transcripts -> chunks with start timestamps (windowed ~60s)."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    entries: list[tuple[float, str]] = []
    cur_ts = None
    for line in raw.splitlines():
        line = line.strip()
        if _TS_RE.search(line):
            cur_ts = _ts_to_sec(line.split("-->")[0].strip())
            continue
        if line and cur_ts is not None and not line.isdigit() and line.upper() != "WEBVTT":
            entries.append((cur_ts, line))
            cur_ts = None
    if not entries:
        # plain transcript without timestamps -> fall back to paragraphs
        return ingest_text(path)
    return _windows(entries, path.name, "transcript")


MEDIA_VIDEO = {".mp4", ".mkv", ".mov", ".webm", ".avi"}
MEDIA_AUDIO = {".mp3", ".wav", ".m4a", ".ogg", ".flac"}


def transcribe(path: Path) -> list[tuple[float, str]]:
    """Speech -> [(start_seconds, text)] with faster-whisper (GPU if usable, else CPU int8). Cached per file."""
    import json
    import os

    cache = path.with_name(path.name + ".whisper.json")
    if cache.exists() and cache.stat().st_mtime >= path.stat().st_mtime:
        return [(float(a), b) for a, b in json.loads(cache.read_text(encoding="utf-8"))]
    # Windows: ctranslate2 and scikit-learn each ship an OpenMP runtime; allow both to load (common workaround)
    os.environ.setdefault("KMP_DUPLICATE_LIB_OK", "TRUE")
    os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")
    try:
        from faster_whisper import WhisperModel
    except ImportError as e:  # pragma: no cover - depends on environment
        raise ValueError("Audio/video transcription needs `pip install faster-whisper`; or upload a .vtt/.srt transcript.") from e
    size = os.environ.get("WHISPER_MODEL", "base")
    last: Exception | None = None
    for device, compute in (("cuda", "float16"), ("cpu", "int8")):
        if device == "cuda" and os.environ.get("WHISPER_DEVICE", "auto") == "cpu":
            continue
        try:
            model = WhisperModel(size, device=device, compute_type=compute)
            segments, _info = model.transcribe(str(path), vad_filter=True)
            out = [(float(sg.start), sg.text.strip()) for sg in segments if sg.text.strip()]
            break
        except Exception as e:  # missing CUDA libs, download failure, ...
            last, out = e, []
    else:
        raise ValueError(f"Transcription failed ({last}). Upload a .vtt/.srt transcript instead.")
    cache.write_text(json.dumps(out, ensure_ascii=False), encoding="utf-8")
    return out


def ingest_media(path: Path) -> list[Unit]:
    """Lecture video/audio -> timestamped transcript units (no manual preprocessing)."""
    entries = transcribe(path)
    if not entries:
        raise ValueError("No speech was detected in this recording.")
    kind = "video" if path.suffix.lower() in MEDIA_VIDEO else "audio"
    return _windows(entries, path.name, kind, window=30.0)


def ingest_file(path: Path, figure_dir: Path | None = None) -> list[Unit]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return ingest_pdf(path, figure_dir)
    if suffix == ".pptx":
        return ingest_pptx(path, figure_dir)
    if suffix in (".vtt", ".srt"):
        return ingest_transcript(path)
    if suffix in MEDIA_VIDEO | MEDIA_AUDIO:
        return ingest_media(path)
    if suffix in (".txt", ".md"):
        # heuristic: timestamped transcript?
        head = path.read_text(encoding="utf-8", errors="replace")[:500]
        if "-->" in head or head.startswith("WEBVTT"):
            return ingest_transcript(path)
        return ingest_text(path)
    raise ValueError(
        f"Unsupported file type: {suffix} (supported: pdf, pptx, txt, md, vtt, srt, mp4/mkv/mov/webm, mp3/wav/m4a)"
    )


class KnowledgeBase:
    """Collection of units + topic tags, persisted as JSON."""

    def __init__(self, units: list[Unit] | None = None):
        self.units: list[Unit] = units or []

    def add_file(self, path: Path, figure_dir: Path | None = None) -> int:
        new = ingest_file(path, figure_dir)
        existing = {(u.source_file, u.location) for u in self.units}
        added = [u for u in new if (u.source_file, u.location) not in existing]
        self.units.extend(added)
        return len(added)

    def to_json(self) -> str:
        return json.dumps([u.to_dict() for u in self.units], ensure_ascii=False, indent=1)

    @classmethod
    def from_json(cls, raw: str) -> "KnowledgeBase":
        fields = set(Unit.__dataclass_fields__)
        return cls([Unit(**{k: v for k, v in d.items() if k in fields}) for d in json.loads(raw)])

    def save(self, path: Path) -> None:
        path.write_text(self.to_json(), encoding="utf-8")

    @classmethod
    def load(cls, path: Path) -> "KnowledgeBase":
        return cls.from_json(path.read_text(encoding="utf-8"))

    def summary(self) -> dict:
        files: dict[str, int] = {}
        for u in self.units:
            files[u.source_file] = files.get(u.source_file, 0) + 1
        return {"total_units": len(self.units), "files": files}
