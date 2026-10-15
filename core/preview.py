"""Source previews for the Source Viewer: the exact PDF page or PPTX slide as an image.

PDF pages are rendered with PyMuPDF. PowerPoint cannot be rendered without Office/LibreOffice, so a slide is
redrawn from its own content (text boxes + pictures at their real positions) with Pillow — a faithful
"slide preview", labelled as such in the UI.
"""
from __future__ import annotations

import io
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

W, H = 1280, 720
_FONTS = ["arial.ttf", "segoeui.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"]
_BOLD = ["arialbd.ttf", "segoeuib.ttf", "DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf"]


def _font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    for name in _BOLD if bold else _FONTS:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def pdf_page_png(pdf: Path, page: int, out_dir: Path, zoom: float = 1.6) -> Path:
    import fitz

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{pdf.stem}-p{page}.png"
    if not out.exists() or out.stat().st_mtime < pdf.stat().st_mtime:
        with fitz.open(pdf) as doc:
            if not 1 <= page <= doc.page_count:
                raise ValueError("page out of range")
            doc[page - 1].get_pixmap(matrix=fitz.Matrix(zoom, zoom)).save(out)
    return out


def slide_png(pptx: Path, slide_no: int, out_dir: Path) -> Path:
    from pptx import Presentation

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{pptx.stem}-s{slide_no}.png"
    if out.exists() and out.stat().st_mtime >= pptx.stat().st_mtime:
        return out
    prs = Presentation(str(pptx))
    if not 1 <= slide_no <= len(prs.slides):
        raise ValueError("slide out of range")
    slide = prs.slides[slide_no - 1]
    sw, sh = prs.slide_width, prs.slide_height
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, 10], fill=(60, 90, 160))
    first_text = True
    for shape in slide.shapes:
        x, y = int(shape.left / sw * W), int(shape.top / sh * H)
        w, h = int(shape.width / sw * W), int(shape.height / sh * H)
        if shape.shape_type == 13:
            try:
                pic = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
                pic.thumbnail((w, h))
                img.paste(pic, (x, y))
            except Exception:
                d.rectangle([x, y, x + w, y + h], outline=(180, 180, 180))
            continue
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            if not text:
                continue
            size = int((para.font.size.pt if para.font.size else 20) * 1.55)
            f = _font(size, bold=bool(para.font.bold) or (first_text and size > 40))
            chars = max(12, int(w / (size * 0.52)))
            for line in textwrap.wrap(text, chars) or [text]:
                d.text((x, y), line, fill=(25, 25, 25), font=f)
                y += int(size * 1.3)
            y += int(size * 0.35)
            first_text = False
    img.save(out)
    return out
